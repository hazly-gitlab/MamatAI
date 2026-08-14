import os
import csv
from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl

def extract_text_from_pdf(file_path: str) -> str:
    try:
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text_parts.append(content)
        return "\n".join(text_parts)
    except Exception as e:
        raise ValueError(f"Error reading PDF: {str(e)}")

def extract_text_from_docx(file_path: str) -> str:
    try:
        doc = DocxDocument(file_path)
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        raise ValueError(f"Error reading DOCX: {str(e)}")

def extract_text_from_xlsx(file_path: str) -> str:
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            text_parts.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                if any(row):
                    row_str = " | ".join([str(val) if val is not None else "" for val in row])
                    text_parts.append(row_str)
        return "\n".join(text_parts)
    except Exception as e:
        raise ValueError(f"Error reading XLSX: {str(e)}")

def extract_text_from_csv(file_path: str) -> str:
    try:
        text_parts = []
        with open(file_path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for row in reader:
                if row:
                    text_parts.append(", ".join(row))
        return "\n".join(text_parts)
    except Exception as e:
        raise ValueError(f"Error reading CSV: {str(e)}")

def extract_text_from_txt(file_path: str) -> str:
    try:
        with open(file_path, mode="r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        raise ValueError(f"Error reading text file: {str(e)}")

def extract_document_text(file_path: str, content_type: str) -> str:
    _, ext = os.path.splitext(file_path.lower())

    if ext == ".pdf" or "pdf" in content_type:
        return extract_text_from_pdf(file_path)
    elif ext == ".docx" or "officedocument.wordprocessingml" in content_type:
        return extract_text_from_docx(file_path)
    elif ext in [".xlsx", ".xls"] or "officedocument.spreadsheetml" in content_type:
        return extract_text_from_xlsx(file_path)
    elif ext == ".csv" or "csv" in content_type:
        return extract_text_from_csv(file_path)
    elif ext in [".txt", ".md", ".markdown"] or "text/" in content_type:
        return extract_text_from_txt(file_path)
    else:
        # Fallback to general text extraction
        return extract_text_from_txt(file_path)
